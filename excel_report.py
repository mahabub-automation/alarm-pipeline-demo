"""
Excel report builder.

Turns the processor output into a formatted multi-sheet workbook:
Summary, Regions, Long Outages and Top Sites.
"""

import os

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from processor import LONG_OUTAGE_MINUTES, format_minutes, process

OUTPUT_DIR = "output"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
LABEL_FONT = Font(bold=True)
THIN_SIDE = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

SEVERITY_FILLS = {
    "Critical": PatternFill("solid", fgColor="F8CBAD"),
    "Major": PatternFill("solid", fgColor="FFE699"),
    "Minor": PatternFill("solid", fgColor="E2EFDA"),
}


def autosize_columns(sheet, max_width=50):
    """Set each column width from its longest value."""
    for column_cells in sheet.columns:
        longest = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[letter].width = min(longest + 6, max_width)


def style_table(sheet):
    """Header colours, borders, frozen header row, filter and column widths."""
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    autosize_columns(sheet)


def color_rows_by_severity(sheet, severity_column):
    """Fill each data row with the colour of its severity."""
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        fill = SEVERITY_FILLS.get(row[severity_column - 1].value)
        if fill:
            for cell in row:
                cell.fill = fill


def color_regions(sheet):
    """Heat-map the Critical and LongOutages columns."""
    last_row = sheet.max_row
    sheet.conditional_formatting.add(
        f"B2:B{last_row}",
        ColorScaleRule(start_type="min", start_color="FFFFFF",
                       end_type="max", end_color="F8696B"),
    )
    sheet.conditional_formatting.add(
        f"F2:F{last_row}",
        ColorScaleRule(start_type="min", start_color="FFFFFF",
                       end_type="max", end_color="F4B183"),
    )


def write_block(sheet, start_row, headers, rows):
    """Write a small styled table and return the next free row."""
    for column, text in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=column, value=text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER

    for offset, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=start_row + offset, column=column, value=value)
            cell.number_format = "#,##0"
            cell.border = BORDER

    return start_row + len(rows) + 2


def write_summary(workbook, result):
    """Create the Summary sheet as the first tab."""
    sheet = workbook.create_sheet("Summary", 0)

    sheet["A1"] = "Telecom Alarm Report"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = "Generated at"
    sheet["A2"].font = LABEL_FONT
    sheet["B2"] = result["generated_at"].strftime("%Y-%m-%d %H:%M")

    active = result["active"]
    metrics = [
        ("Raw records", int(result["raw_count"])),
        ("Clean records", int(result["clean_count"])),
        ("Active alarms", int(len(active))),
        ("Active critical alarms", int((active["Severity"] == "Critical").sum())),
        (f"Long outages (>= {LONG_OUTAGE_MINUTES // 60}h)", int(len(result["long_outages"]))),
        ("Sites with active alarms", int(active["SiteID"].nunique())),
    ]
    next_row = write_block(sheet, 4, ["Metric", "Value"], metrics)

    issues = [
        (name.replace("_", " ").capitalize(), int(value))
        for name, value in result["issues"].items()
    ]
    write_block(sheet, next_row, ["Data quality check", "Rows found"], issues)

    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 20


def long_outage_frame(result):
    """Columns shown on the Long Outages sheet."""
    frame = result["long_outages"][
        ["SiteID", "Region", "AlarmName", "Severity", "OccurTime", "DurationMin"]
    ].copy()
    frame["Duration"] = frame["DurationMin"].apply(format_minutes)
    return frame[
        ["SiteID", "Region", "AlarmName", "Severity", "OccurTime", "Duration", "DurationMin"]
    ]


def top_sites_frame(result):
    """Columns shown on the Top Sites sheet."""
    frame = result["top_sites"].copy()
    frame["Longest"] = frame["LongestMin"].apply(format_minutes)
    return frame.drop(columns="LongestMin")


def build_report(result, output_dir=OUTPUT_DIR):
    """Write the workbook and return its path."""
    os.makedirs(output_dir, exist_ok=True)
    stamp = result["generated_at"].strftime("%Y%m%d_%H%M")
    path = os.path.join(output_dir, f"alarm_report_{stamp}.xlsx")

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        result["regions"].to_excel(writer, sheet_name="Regions", index=False)
        long_outage_frame(result).to_excel(writer, sheet_name="Long Outages", index=False)
        top_sites_frame(result).to_excel(writer, sheet_name="Top Sites", index=False)

        book = writer.book
        for name in ["Regions", "Long Outages", "Top Sites"]:
            style_table(book[name])

        color_regions(book["Regions"])

        long_sheet = book["Long Outages"]
        color_rows_by_severity(long_sheet, severity_column=4)
        for cell in long_sheet["E"][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"

        write_summary(book, result)

        # Open on the Summary tab, with only one tab selected
        for sheet in book.worksheets:
            sheet.sheet_view.tabSelected = False
        book.active = 0
        book["Summary"].sheet_view.tabSelected = True

    return path


def main():
    result = process()
    path = build_report(result)

    print("=" * 50)
    print("Excel report created")
    print("=" * 50)
    print(f"File          : {path}")
    print("Sheets        : Summary, Regions, Long Outages, Top Sites")
    print(f"Active alarms : {len(result['active']):,}")
    print(f"Long outages  : {len(result['long_outages']):,}")


if __name__ == "__main__":
    main()
